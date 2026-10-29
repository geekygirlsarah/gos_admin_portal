"""Shared spreadsheet generation utilities for CSV and Excel (.xlsx) exports."""

from __future__ import annotations

import csv
from io import BytesIO
from typing import Any, Iterable, Sequence

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def export_csv_response(
    filename: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
) -> HttpResponse:
    """Generate a downloadable CSV HttpResponse from headers and row data."""
    if not filename.lower().endswith(".csv"):
        filename = f"{filename}.csv"

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'

    writer = csv.writer(response)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)

    return response


def _apply_header_styles(ws, font: Font | None, fill: PatternFill | None) -> None:
    if not (font or fill):
        return
    for cell in ws[1]:
        if font:
            cell.font = font
        if fill:
            cell.fill = fill


def _apply_auto_fit_columns(ws, min_col_width: int) -> None:
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 2, min_col_width)


def _apply_currency_formatting(
    ws, currency_columns: Sequence[int], max_col: int
) -> None:
    for row in ws.iter_rows(min_row=2, max_col=max_col):
        for cell in row:
            if cell.column in currency_columns and cell.value not in ("", None):
                cell.number_format = "$#,##0.00"


def export_xlsx_response(
    filename: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
    sheet_title: str = "Sheet1",
    header_font: Font | None = None,
    header_fill: PatternFill | None = None,
    auto_fit_columns: bool = False,
    min_col_width: int = 12,
    currency_columns: Sequence[int] | None = None,
    customizer: Any = None,
) -> HttpResponse:
    """Generate a downloadable Excel (.xlsx) HttpResponse using openpyxl."""
    if not filename.lower().endswith(".xlsx"):
        filename = f"{filename}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.append(list(headers))
    _apply_header_styles(ws, header_font, header_fill)

    for row in rows:
        ws.append(list(row))

    if auto_fit_columns:
        _apply_auto_fit_columns(ws, min_col_width)

    if currency_columns:
        _apply_currency_formatting(ws, currency_columns, max_col=len(headers))

    if customizer:
        customizer(wb, ws)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
