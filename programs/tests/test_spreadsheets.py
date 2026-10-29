"""Tests for shared spreadsheet generation utilities."""

import csv
from io import BytesIO

from django.test import SimpleTestCase
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from programs.utils.spreadsheets import export_csv_response, export_xlsx_response


class SpreadsheetExportTests(SimpleTestCase):
    def test_export_csv_response(self):
        headers = ["Name", "Role", "Active"]
        rows = [
            ["Alice Smith", "Student", True],
            ["Bob Jones", "Mentor", False],
        ]
        response = export_csv_response("test_export.csv", headers, rows)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")
        self.assertEqual(
            response["Content-Disposition"],
            'attachment; filename="test_export.csv"',
        )

        content = response.content.decode("utf-8")
        reader = list(csv.reader(content.splitlines()))
        self.assertEqual(reader[0], ["Name", "Role", "Active"])
        self.assertEqual(reader[1], ["Alice Smith", "Student", "True"])
        self.assertEqual(reader[2], ["Bob Jones", "Mentor", "False"])

    def test_export_xlsx_response_default_styling(self):
        headers = ["Item", "Qty", "Price"]
        rows = [
            ["Robotics Kit", 2, 49.99],
            ["Screws", 100, 0.05],
        ]
        response = export_xlsx_response(
            "inventory.xlsx", headers, rows, sheet_title="Inventory"
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            response["Content-Disposition"],
            'attachment; filename="inventory.xlsx"',
        )

        wb = load_workbook(BytesIO(response.content))
        self.assertIn("Inventory", wb.sheetnames)
        ws = wb["Inventory"]

        cell_headers = [c.value for c in ws[1]]
        self.assertEqual(cell_headers, ["Item", "Qty", "Price"])

        row_1 = [c.value for c in ws[2]]
        self.assertEqual(row_1, ["Robotics Kit", 2, 49.99])
        row_2 = [c.value for c in ws[3]]
        self.assertEqual(row_2, ["Screws", 100, 0.05])

    def test_export_xlsx_response_custom_styling(self):
        headers = ["Header 1", "Header 2"]
        rows = [["Val 1", "Val 2"]]
        custom_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
        custom_fill = PatternFill("solid", fgColor="003366")

        response = export_xlsx_response(
            "custom.xlsx",
            headers,
            rows,
            sheet_title="CustomSheet",
            header_font=custom_font,
            header_fill=custom_fill,
        )

        wb = load_workbook(BytesIO(response.content))
        ws = wb["CustomSheet"]
        self.assertEqual(ws["A1"].font.bold, True)
        self.assertEqual(ws["A1"].fill.fill_type, "solid")
