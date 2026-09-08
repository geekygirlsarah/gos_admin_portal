import csv
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from orders.forms import OrderItemForm
from orders.models import Order, OrderItem
from orders.tests.base import (
    make_item,
    make_lead_mentor_user,
    make_mentor_user,
    make_order,
    make_program,
    make_student_user,
)
from programs.models import Crew, Program, SubTeam, Team


class OrderItemStudentLabelsTests(TestCase):
    def setUp(self):
        self.program = make_program(name="Main Program")
        self.other_program = make_program(name="Other Program")
        self.lead = make_lead_mentor_user()
        self.mentor = make_mentor_user()
        self.student = make_student_user(program=self.program)

        self.team1 = Team.objects.create(
            team_type="FRC", number=3504, name="Girls of Steel", color="#dc3545"
        )
        self.team2 = Team.objects.create(
            team_type="FTC", number=9820, name="Junior Team", color="#0d6efd"
        )

        self.crew1 = Crew.objects.create(
            program=self.program, name="Chassis Crew", color="#198754"
        )
        self.crew_other = Crew.objects.create(
            program=self.other_program, name="Other Crew", color="#ffc107"
        )

        self.subteam1 = SubTeam.objects.create(
            program=self.program, name="Mechanical", color="#6f42c1"
        )
        self.subteam_other = SubTeam.objects.create(
            program=self.other_program, name="Other Subteam", color="#fd7e14"
        )

        self.list_url = reverse("orders:order_list", args=[self.program.id])

    def login(self, user):
        self.client.force_login(user)

    def test_form_filters_crews_and_subteams_by_program(self):
        form = OrderItemForm(program=self.program)
        self.assertIn(self.team1, form.fields["team"].queryset)
        self.assertIn(self.team2, form.fields["team"].queryset)
        self.assertIn(self.crew1, form.fields["crew"].queryset)
        self.assertNotIn(self.crew_other, form.fields["crew"].queryset)
        self.assertIn(self.subteam1, form.fields["subteam"].queryset)
        self.assertNotIn(self.subteam_other, form.fields["subteam"].queryset)

    def test_student_requests_item_with_labels(self):
        self.login(self.student)
        create_url = reverse("orders:item_create", args=[self.program.id])
        resp = self.client.post(
            create_url,
            {
                "item_name": "Bearing Kit",
                "quantity": "2",
                "unit_price": "12.50",
                "team": str(self.team1.pk),
                "crew": str(self.crew1.pk),
                "subteam": str(self.subteam1.pk),
                "vendor_choice": "",
            },
        )
        self.assertRedirects(resp, self.list_url)
        item = OrderItem.objects.get(item_name="Bearing Kit")
        self.assertEqual(item.team, self.team1)
        self.assertEqual(item.crew, self.crew1)
        self.assertEqual(item.subteam, self.subteam1)

    def test_item_without_labels_allowed(self):
        self.login(self.student)
        create_url = reverse("orders:item_create", args=[self.program.id])
        resp = self.client.post(
            create_url,
            {
                "item_name": "Screws",
                "quantity": "1",
                "team": "",
                "crew": "",
                "subteam": "",
                "vendor_choice": "",
            },
        )
        self.assertRedirects(resp, self.list_url)
        item = OrderItem.objects.get(item_name="Screws")
        self.assertIsNone(item.team)
        self.assertIsNone(item.crew)
        self.assertIsNone(item.subteam)

    def test_order_list_renders_label_badges(self):
        make_item(
            self.program,
            item_name="Bearing Kit",
            team=self.team1,
            crew=self.crew1,
            subteam=self.subteam1,
        )
        self.login(self.lead)
        resp = self.client.get(self.list_url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "FRC 3504 Girls of Steel")
        self.assertContains(resp, "Chassis Crew")
        self.assertContains(resp, "Mechanical")

    def test_order_detail_renders_label_badges(self):
        order = make_order(self.program, created_by=self.lead)
        item = make_item(
            self.program,
            item_name="Motor Controller",
            order=order,
            team=self.team1,
            crew=self.crew1,
            subteam=self.subteam1,
        )
        self.login(self.lead)
        detail_url = reverse("orders:order_detail", args=[self.program.id, order.pk])
        resp = self.client.get(detail_url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "FRC 3504 Girls of Steel")
        self.assertContains(resp, "Chassis Crew")
        self.assertContains(resp, "Mechanical")

    def test_order_archive_renders_label_badges(self):
        order = make_order(
            self.program, created_by=self.lead, status=Order.STATUS_RECEIVED
        )
        item = make_item(
            self.program,
            item_name="Sensor Board",
            order=order,
            team=self.team1,
            crew=self.crew1,
        )
        self.login(self.lead)
        archive_url = reverse("orders:order_archive", args=[self.program.id])
        resp = self.client.get(archive_url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "FRC 3504 Girls of Steel")
        self.assertContains(resp, "Chassis Crew")

    def test_csv_export_includes_team_crew_subteam(self):
        make_item(
            self.program,
            item_name="Pneumatics Valve",
            team=self.team1,
            crew=self.crew1,
            subteam=self.subteam1,
        )
        self.login(self.lead)
        resp = self.client.get(f"{self.list_url}?export=csv")
        self.assertEqual(resp.status_code, 200)
        content = resp.content.decode("utf-8")
        reader = csv.reader(content.splitlines())
        rows = list(reader)
        headers = rows[0]
        self.assertIn("Team", headers)
        self.assertIn("Crew", headers)
        self.assertIn("Subteam", headers)
        self.assertNotIn("Tag", headers)

        data_row = rows[1]
        team_idx = headers.index("Team")
        crew_idx = headers.index("Crew")
        subteam_idx = headers.index("Subteam")
        self.assertEqual(data_row[team_idx], str(self.team1))
        self.assertEqual(data_row[crew_idx], self.crew1.name)
        self.assertEqual(data_row[subteam_idx], self.subteam1.name)

    def test_xlsx_export_includes_team_crew_subteam(self):
        make_item(
            self.program,
            item_name="Pneumatics Valve",
            team=self.team1,
            crew=self.crew1,
            subteam=self.subteam1,
        )
        self.login(self.lead)
        resp = self.client.get(f"{self.list_url}?export=xlsx")
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Team", headers)
        self.assertIn("Crew", headers)
        self.assertIn("Subteam", headers)
        self.assertNotIn("Tag", headers)

    def test_deleting_team_crew_subteam_sets_null_on_item(self):
        item = make_item(
            self.program,
            item_name="Omni Wheel",
            team=self.team1,
            crew=self.crew1,
            subteam=self.subteam1,
        )
        self.team1.delete()
        self.crew1.delete()
        self.subteam1.delete()
        item.refresh_from_db()
        self.assertIsNone(item.team)
        self.assertIsNone(item.crew)
        self.assertIsNone(item.subteam)
        self.assertTrue(OrderItem.objects.filter(pk=item.pk).exists())
